#!/usr/bin/env python3
"""Rebuild Member 4 analytics from measured, provenance-checked benchmark outputs."""

import csv
import json
import math
import os
import statistics
from collections import defaultdict

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from scipy.stats import rankdata, wilcoxon

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SCRIPT_DIR)
RESULTS = os.path.join(ROOT, "results")
MASTER_CSV = os.path.join(RESULTS, "master_benchmark_summary.csv")
CATALOG_JSON = os.path.join(ROOT, "target_benchmark", "all_bugs_catalog.json")
MIO_CSV = os.path.join(ROOT, "MIO_Algorithm", "Result_Round2", "evosuite_budget_summary.csv")
AI_CSV = os.path.join(RESULTS, "Deepseek_vs_Gemini_Economics.csv")
TECHNIQUES = ["IPO (Native IPO)", "MIO (EvoSuite SBST)", "DeepSeek V4 Flash", "Gemini 3.8 Flash"]
ATTEMPTED = {"DONE", "COMPILE_ERROR", "TIMEOUT"}
UNRESOLVED = {"NOT_RUN", "STALE_RESULT", "CHECKOUT_ERROR", "INVALID_SUITE", "RUN_ERROR"}


def number(value):
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def read_csv(path):
    if not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as stream:
        return list(csv.DictReader(stream))


def load_master():
    rows = read_csv(MASTER_CSV)
    for row in rows:
        try:
            row["Bug_ID"] = int(row["Bug_ID"])
        except (TypeError, ValueError):
            row["Bug_ID"] = None
        for field in ("Line_Coverage_%", "Branch_Coverage_%", "Evaluation_Duration_Sec"):
            row[field] = number(row.get(field))
    return rows


def load_catalog():
    with open(CATALOG_JSON, encoding="utf-8") as stream:
        catalog = json.load(stream)
    bugs = {(item["project"], int(item["bug_id"])) for item in catalog}
    single, multi = set(), set()
    for item in catalog:
        key = (item["project"], int(item["bug_id"]))
        classes = item.get("target_classes") or item.get("modified_classes") or []
        (multi if len(classes) > 1 else single).add(key)
    return bugs, single, multi


def summarize(values):
    return {
        "n": len(values),
        "mean": round(statistics.mean(values), 3) if values else None,
        "sd": round(statistics.stdev(values), 3) if len(values) > 1 else (0.0 if values else None),
    }


def paired_rank_biserial(differences):
    nonzero = [value for value in differences if value != 0]
    if not nonzero:
        return 0.0
    ranks = rankdata(np.abs(nonzero))
    positive = sum(rank for rank, value in zip(ranks, nonzero) if value > 0)
    negative = sum(rank for rank, value in zip(ranks, nonzero) if value < 0)
    total = positive + negative
    return round(float((positive - negative) / total), 4) if total else 0.0


def holm_adjust(results):
    measured = [(i, item["p_value"]) for i, item in enumerate(results) if item.get("p_value") is not None]
    ordered = sorted(measured, key=lambda item: item[1])
    running = 0.0
    count = len(ordered)
    for rank, (index, p_value) in enumerate(ordered):
        running = max(running, min(1.0, (count - rank) * p_value))
        results[index]["p_value_holm"] = running
    return results


def run_hypothesis(rows):
    valid = {
        tech: {(r["Project"], r["Bug_ID"]): r["Line_Coverage_%"] for r in rows
               if r["Technique"] == tech and r["Execution_Status"] == "DONE"
               and r["Line_Coverage_%"] is not None}
        for tech in TECHNIQUES
    }
    pairs = [
        ("Gemini 3.8 Flash", "DeepSeek V4 Flash"),
        ("MIO (EvoSuite SBST)", "Gemini 3.8 Flash"),
        ("MIO (EvoSuite SBST)", "DeepSeek V4 Flash"),
        ("MIO (EvoSuite SBST)", "IPO (Native IPO)"),
        ("Gemini 3.8 Flash", "IPO (Native IPO)"),
        ("DeepSeek V4 Flash", "IPO (Native IPO)"),
    ]
    results = []
    for left, right in pairs:
        common = sorted(set(valid[left]) & set(valid[right]))
        x = [valid[left][key] for key in common]
        y = [valid[right][key] for key in common]
        differences = [a - b for a, b in zip(x, y)]
        item = {"group_1": left, "group_2": right, "n_pairs": len(common),
                "n_nonzero_pairs": sum(value != 0 for value in differences)}
        if len(common) < 2:
            item["status"] = "INSUFFICIENT_SAMPLE"
        else:
            if item["n_nonzero_pairs"] == 0:
                stat, p_value = 0.0, 1.0
            else:
                stat, p_value = wilcoxon(x, y, alternative="two-sided", zero_method="wilcox")
            item.update({"status": "MEASURED", "wilcoxon_statistic": float(stat),
                         "p_value": float(p_value), "paired_rank_biserial": paired_rank_biserial(differences)})
        results.append(item)
    return holm_adjust(results)


def run_mio_budget():
    records = read_csv(MIO_CSV)
    grouped = defaultdict(lambda: {"line": [], "branch": [], "duration": []})
    by_budget_key = defaultdict(dict)
    for row in records:
        budget = number(row.get("Budget_Sec"))
        if budget is None:
            continue
        budget = int(budget)
        for key, source in (("line", "Line_Cov_Mean_%"), ("branch", "Branch_Cov_Mean_%"),
                            ("duration", "Avg_Duration_Sec")):
            value = number(row.get(source))
            if value is not None:
                grouped[int(budget)][key].append(value)
                if key == "line":
                    target_key = (row.get("Project", ""), row.get("Bug_ID", ""), row.get("Target_Class", ""))
                    by_budget_key[budget][target_key] = value
    summary = {}
    for budget in (30, 60, 120):
        vals = grouped[budget]
        summary[budget] = {
            "eval_count": len(vals["line"]),
            "line": summarize(vals["line"]),
            "branch": summarize(vals["branch"]),
            "duration": summarize(vals["duration"]),
        }
    tests = {}
    for left, right in ((30, 60), (60, 120)):
        common = sorted(set(by_budget_key[left]) & set(by_budget_key[right]))
        x = [by_budget_key[left][key] for key in common]
        y = [by_budget_key[right][key] for key in common]
        differences = [a - b for a, b in zip(x, y)]
        if len(common) > 1:
            if all(value == 0 for value in differences):
                stat, p_value = 0.0, 1.0
            else:
                stat, p_value = wilcoxon(x, y, alternative="two-sided", zero_method="wilcox")
            tests[f"{left}s_vs_{right}s"] = {
                "status": "MEASURED", "n_paired_target_classes": len(common),
                "n_nonzero_pairs": sum(value != 0 for value in differences),
                "wilcoxon_statistic": float(stat), "p_value": float(p_value),
                "paired_rank_biserial": paired_rank_biserial(differences),
            }
        else:
            tests[f"{left}s_vs_{right}s"] = {"status": "INSUFFICIENT_SAMPLE"}
    test_rows = [dict(value, comparison=key) for key, value in tests.items() if value.get("p_value") is not None]
    holm_adjust(test_rows)
    for row in test_rows:
        tests[row["comparison"]]["p_value_holm"] = row.get("p_value_holm")
    deltas = {}
    for left, right in ((30, 60), (60, 120)):
        a, b = summary[left]["line"]["mean"], summary[right]["line"]["mean"]
        deltas[f"gain_{left}_to_{right}"] = round(b - a, 3) if a is not None and b is not None else None
    result = {"source": "MIO generation budget summary; separate from benchmark evaluations",
              "budget_summary": summary, "delta_coverage": deltas, "paired_tests": tests}
    available = [b for b in (30, 60, 120) if summary[b]["line"]["mean"] is not None]
    fig, axes = plt.subplots(1, 2, figsize=(12, 5))
    if not available:
        for ax in axes:
            ax.axis("off")
            ax.text(.5, .5, "No measured MIO budget records", ha="center", va="center")
    else:
        means = [summary[b]["line"]["mean"] for b in available]
        branches = [summary[b]["branch"]["mean"] for b in available]
        axes[0].plot(available, means, marker="o", label="Line coverage")
        branch_pts = [(b, summary[b]["branch"]["mean"]) for b in available if summary[b]["branch"]["mean"] is not None]
        if branch_pts:
            axes[0].plot([p[0] for p in branch_pts], [p[1] for p in branch_pts], marker="s", label="Branch coverage")
        axes[0].set(xlabel="EvoSuite MIO generation budget (seconds)", ylabel="Coverage mean (%)",
                    title="MIO generation coverage by budget")
        axes[0].set_xticks(available, [f"{b}s\n(n={summary[b]['eval_count']})" for b in available])
        axes[0].legend()
        gains = [summary[b]["line"]["mean"] for b in available]
        axes[1].bar([f"{b}s\n(n={summary[b]['eval_count']})" for b in available], gains, color="#2b5c8f")
        axes[1].set(xlabel="Generation budget", ylabel="Line coverage mean (%)", title="Observed line coverage")
    fig.tight_layout()
    fig.savefig(os.path.join(RESULTS, "figure5_budget_scaling.png"), dpi=300)
    plt.close(fig)
    return result


def run_ensemble(rows, catalog_bugs):
    detected, attempted = defaultdict(set), defaultdict(set)
    for row in rows:
        key = (row["Project"], row["Bug_ID"])
        tech = row["Technique"]
        if row["Execution_Status"] in ATTEMPTED:
            attempted[tech].add(key)
        if row["Fault_Detection_Status"] == "BUG_DETECTED":
            detected[tech].add(key)
    union = set().union(*(detected[t] for t in TECHNIQUES))
    any_attempted = set().union(*(attempted[t] for t in TECHNIQUES))
    matrix_complete = not any(r["Execution_Status"] in UNRESOLVED for r in rows)
    available_suites_complete = all(
        r["Execution_Status"] in ATTEMPTED for r in rows if r.get("Suite_Available") == "YES"
    )
    result = {
        "catalog_bugs": len(catalog_bugs),
        "results_complete": matrix_complete and available_suites_complete,
        "available_suite_evaluations_complete": available_suites_complete,
        "bugs_with_any_attempt": len(any_attempted),
        "ensemble_detected_bugs": len(union),
        "ensemble_fdr_full_catalog_percent": round(100 * len(union) / len(catalog_bugs), 3) if catalog_bugs else None,
        "ensemble_fdr_among_any_attempt_percent": round(100 * len(union) / len(any_attempted), 3) if any_attempted else None,
        "attempted_per_technique": {t: len(attempted[t]) for t in TECHNIQUES},
        "detected_per_technique": {t: len(detected[t]) for t in TECHNIQUES},
        "unique_contributions": {t: len(detected[t] - set().union(*(detected[x] for x in TECHNIQUES if x != t)))
                                 for t in TECHNIQUES},
        "overlaps": {},
    }
    for i, left in enumerate(TECHNIQUES):
        for right in TECHNIQUES[i + 1:]:
            result["overlaps"][f"{left} & {right}"] = len(detected[left] & detected[right])
    fig, axes = plt.subplots(1, 2, figsize=(13, 5.5))
    counts = [len(detected[t]) for t in TECHNIQUES] + [len(union)]
    labels = [f"{name}\n(n={len(attempted[tech])})" for name, tech in zip(
        ("IPO", "MIO", "DeepSeek", "Gemini"), TECHNIQUES
    )] + ["Ensemble"]
    bars = axes[0].bar(labels, counts, color=["#2ca02c", "#888888", "#4a90e2", "#1a73e8", "#d9534f"])
    axes[0].set(ylabel="Bugs detected", title="Measured bug detections")
    axes[0].set_ylim(0, max(1, max(counts, default=0) * 1.3))
    for bar in bars:
        axes[0].text(bar.get_x() + bar.get_width()/2, bar.get_height(), f"{int(bar.get_height())}",
                     ha="center", va="bottom")
    uniq = [result["unique_contributions"][t] for t in TECHNIQUES]
    axes[1].bar(labels[:4], uniq, color="#2ca02c")
    axes[1].set(ylabel="Bugs detected by this technique only", title="Unique contribution")
    axes[1].set_ylim(0, max(1, max(uniq, default=0) * 1.3))
    snapshot = "complete" if result["results_complete"] else "provisional snapshot"
    fig.suptitle(f"Measured Bug Detection ({snapshot}; catalog={len(catalog_bugs)})", fontsize=14)
    fig.tight_layout()
    fig.savefig(os.path.join(RESULTS, "figure6_ensemble_overlap.png"), dpi=300)
    plt.close(fig)
    return result


def run_single_multi(rows, single_bugs, multi_bugs):
    out = {}
    for tech in TECHNIQUES:
        group_result = {}
        for name, keys in (("single_class", single_bugs), ("multi_class", multi_bugs)):
            selected = [r for r in rows if r["Technique"] == tech and (r["Project"], r["Bug_ID"]) in keys]
            attempts = [r for r in selected if r["Execution_Status"] in ATTEMPTED]
            cov = [r["Line_Coverage_%"] for r in selected
                   if r["Execution_Status"] == "DONE" and r["Line_Coverage_%"] is not None]
            detected = sum(r["Fault_Detection_Status"] == "BUG_DETECTED" for r in attempts)
            group_result[name] = {
                "catalog_bugs": len(keys), "attempted_bugs": len(attempts),
                "coverage_n": len(cov), "line_coverage_mean": summarize(cov)["mean"],
                "detected_bugs": detected,
                "fdr_attempted_percent": round(100 * detected / len(attempts), 3) if attempts else None,
            }
        out[tech] = group_result
    return out


def run_ai_economics(rows):
    source = read_csv(AI_CSV)
    grouped = defaultdict(lambda: {"tokens": [], "latency": [], "records": 0})
    for row in source:
        model = (row.get("Model") or "").lower()
        key = "Gemini 3.8 Flash" if "gemini" in model else "DeepSeek V4 Flash" if "deepseek" in model else None
        if not key:
            continue
        tokens = number(row.get("Total_Tokens"))
        latency = number(row.get("Generation_Time_Sec") or row.get("Generation_Time_s"))
        grouped[key]["records"] += 1
        if tokens is not None:
            grouped[key]["tokens"].append(tokens)
        if latency is not None:
            grouped[key]["latency"].append(latency)
    result = {}
    for tech in ("Gemini 3.8 Flash", "DeepSeek V4 Flash"):
        detected = { (r["Project"], r["Bug_ID"]) for r in rows
                     if r["Technique"] == tech and r["Fault_Detection_Status"] == "BUG_DETECTED" }
        g = grouped[tech]
        result[tech] = {
            "source": os.path.relpath(AI_CSV, ROOT),
            "generation_records": g["records"],
            "avg_tokens_per_generation_record": round(statistics.mean(g["tokens"]), 2) if g["tokens"] else None,
            "avg_generation_latency_sec": round(statistics.mean(g["latency"]), 2) if g["latency"] else None,
            "benchmark_detected_bugs": len(detected),
            "token_cost_per_detected_bug": None,
            "note": "Generation logs are summarized separately; they do not carry benchmark run IDs, so token cost per measured detection is not joined.",
        }
    return result


def export_outputs(rows, catalog_bugs, hypothesis, budget, ensemble, single_multi, economics):
    available_suite_rows = sum(r.get("Suite_Available") == "YES" for r in rows)
    attempted_rows = sum(r["Execution_Status"] in ATTEMPTED for r in rows)
    no_suite_rows = sum(r["Execution_Status"] == "NO_SUITE" for r in rows)
    unresolved_rows = sum(r["Execution_Status"] in UNRESOLVED for r in rows)
    results_complete = unresolved_rows == 0 and all(
        r["Execution_Status"] in ATTEMPTED for r in rows if r.get("Suite_Available") == "YES"
    )
    stats = {
        "dataset": {
            "catalog_bugs": len(catalog_bugs), "projects": len({p for p, _ in catalog_bugs}),
            "expected_bug_technique_rows": len(catalog_bugs) * len(TECHNIQUES),
            "master_rows": len(rows),
            "available_suite_rows": available_suite_rows,
            "attempted_rows": attempted_rows,
            "no_suite_rows": no_suite_rows,
            "unresolved_rows": unresolved_rows,
            "results_complete": results_complete,
            "available_suite_evaluations_complete": all(
                r["Execution_Status"] in ATTEMPTED
                for r in rows if r.get("Suite_Available") == "YES"
            ),
            "no_suite_rows": sum(r["Execution_Status"] == "NO_SUITE" for r in rows),
            "unresolved_rows": sum(r["Execution_Status"] in UNRESOLVED for r in rows),
        },
        "hypothesis_testing": hypothesis,
        "mio_generation_budget": budget,
        "ensemble": ensemble,
        "single_vs_multiclass": single_multi,
        "ai_economics": economics,
    }
    with open(os.path.join(RESULTS, "advanced_analytics.json"), "w", encoding="utf-8") as stream:
        json.dump(stats, stream, indent=2, ensure_ascii=False)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    headers = ["Technique", "Catalog bugs", "Suites available", "Attempted", "Completed", "Detected",
               "FDR of attempts (%)", "FDR of catalog (%)", "Line coverage mean (%)",
               "Branch coverage mean (%)", "Mean evaluation seconds"]
    ws.append(["Member 4 benchmark summary — generated from current master CSV"])
    ws.append(headers)
    for cell in ws[2]:
        cell.fill = PatternFill("solid", fgColor="1F497D")
        cell.font = Font(color="FFFFFF", bold=True)
    for tech in TECHNIQUES:
        tr = [r for r in rows if r["Technique"] == tech]
        attempts = [r for r in tr if r["Execution_Status"] in ATTEMPTED]
        done = [r for r in tr if r["Execution_Status"] == "DONE"]
        detected = sum(r["Fault_Detection_Status"] == "BUG_DETECTED" for r in attempts)
        line = [r["Line_Coverage_%"] for r in done if r["Line_Coverage_%"] is not None]
        branch = [r["Branch_Coverage_%"] for r in done if r["Branch_Coverage_%"] is not None]
        dur = [r["Evaluation_Duration_Sec"] for r in attempts if r["Evaluation_Duration_Sec"] is not None]
        suites = sum(r.get("Suite_Available") == "YES" for r in tr)
        ws.append([tech, len(catalog_bugs), suites, len(attempts), len(done), detected,
                   round(100*detected/len(attempts), 3) if attempts else None,
                   round(100*detected/len(catalog_bugs), 3) if catalog_bugs else None,
                   summarize(line)["mean"], summarize(branch)["mean"],
                   summarize(dur)["mean"]])
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:K{ws.max_row}"

    master = wb.create_sheet("Master evaluations")
    if rows:
        fields = list(rows[0].keys())
        master.append(fields)
        for row in rows:
            master.append([row.get(field) for field in fields])
        master.freeze_panes = "A2"
        master.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{master.max_row}"

    mio_sheet = wb.create_sheet("MIO generation budget")
    mio_sheet.append(["Budget (sec)", "Generation records", "Line mean (%)", "Line SD (%)",
                      "Branch mean (%)", "Branch SD (%)", "Mean generation time (sec)"])
    for budget_sec in (30, 60, 120):
        item = budget["budget_summary"][budget_sec]
        mio_sheet.append([budget_sec, item["eval_count"], item["line"]["mean"], item["line"]["sd"],
                          item["branch"]["mean"], item["branch"]["sd"], item["duration"]["mean"]])

    hypo_sheet = wb.create_sheet("Hypothesis tests")
    hypo_sheet.append(["Technique 1", "Technique 2", "Matched bug N", "Nonzero pairs", "Status",
                       "Wilcoxon statistic", "p-value", "Holm-adjusted p", "Paired rank-biserial"])
    for item in hypothesis:
        hypo_sheet.append([item.get("group_1"), item.get("group_2"), item.get("n_pairs"),
                           item.get("n_nonzero_pairs"), item.get("status"), item.get("wilcoxon_statistic"),
                           item.get("p_value"), item.get("p_value_holm"), item.get("paired_rank_biserial")])

    ai_sheet = wb.create_sheet("AI generation logs")
    ai_sheet.append(["Model", "Generation records", "Mean tokens/record", "Mean generation sec",
                     "Measured benchmark detections", "Tokens/detection join"])
    for model, item in economics.items():
        ai_sheet.append([model, item["generation_records"], item["avg_tokens_per_generation_record"],
                         item["avg_generation_latency_sec"], item["benchmark_detected_bugs"],
                         "Not joined: no run ID in generation log"])

    ensemble_sheet = wb.create_sheet("Ensemble")
    ensemble_sheet.append(["Metric", "Value"])
    for name in ("catalog_bugs", "results_complete", "bugs_with_any_attempt", "ensemble_detected_bugs",
                 "ensemble_fdr_full_catalog_percent", "ensemble_fdr_among_any_attempt_percent"):
        ensemble_sheet.append([name, ensemble.get(name)])
    for tech in TECHNIQUES:
        ensemble_sheet.append([f"Attempted: {tech}", ensemble["attempted_per_technique"].get(tech, 0)])
        ensemble_sheet.append([f"Detected: {tech}", ensemble["detected_per_technique"].get(tech, 0)])
        ensemble_sheet.append([f"Unique: {tech}", ensemble["unique_contributions"].get(tech, 0)])

    classes_sheet = wb.create_sheet("Single vs multiclass")
    classes_sheet.append(["Technique", "Bug class group", "Catalog bugs", "Attempted", "Coverage N",
                          "Line mean (%)", "Detected", "FDR of attempts (%)"])
    for tech, groups in single_multi.items():
        for group_name, item in groups.items():
            classes_sheet.append([tech, group_name, item["catalog_bugs"], item["attempted_bugs"],
                                  item["coverage_n"], item["line_coverage_mean"], item["detected_bugs"],
                                  item["fdr_attempted_percent"]])

    dictionary_sheet = wb.create_sheet("Data dictionary")
    dictionary_sheet.append(["Field", "Meaning"])
    dictionary_rows = [
        ("Project, Bug_ID", "Defects4J bug key; one row per bug-technique pair"),
        ("Technique", "IPO Native, MIO, DeepSeek, or Gemini"),
        ("Target_Classes", "Modified target classes in scope"),
        ("Suite_Available, Test_Files", "Current suite inventory"),
        ("Line_Coverage_%, Branch_Coverage_%", "Measured target-class coverage; blank unless measured"),
        ("Fault_Detection_Status", "BUG_DETECTED only if buggy fails and fixed passes"),
        ("Execution_Status", "DONE, COMPILE_ERROR, TIMEOUT, NO_SUITE, NOT_RUN, INVALID_SUITE, or STALE_RESULT"),
        ("Suite_SHA256, Run_ID, Timestamp", "Suite provenance and run identity"),
        ("Run_Log", "Structured per-run JSON record with result, provenance, timing, and error details"),
        ("Evaluation_Duration_Sec", "Wall time for benchmark evaluation"),
    ]
    for item in dictionary_rows:
        dictionary_sheet.append(list(item))

    for sheet in wb.worksheets:
        header_row = 2 if sheet.title == "Summary" else 1
        for cell in sheet[header_row]:
            cell.fill = PatternFill("solid", fgColor="1F497D")
            cell.font = Font(color="FFFFFF", bold=True)
        sheet.freeze_panes = "A3" if sheet.title == "Summary" else "A2"
        if sheet.max_row >= header_row:
            sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for sheet in wb.worksheets:
        for col in sheet.columns:
            width = min(max(max(len(str(c.value or "")) for c in col) + 2, 12), 45)
            sheet.column_dimensions[get_column_letter(col[0].column)].width = width
    wb.save(os.path.join(RESULTS, "Master_Benchmark_Results.xlsx"))

    status_counts = defaultdict(lambda: defaultdict(int))
    for row in rows:
        status_counts[row["Technique"]][row["Execution_Status"]] += 1
    lines = [
        "# Dataset specification",
        "",
        f"Catalog: {len(catalog_bugs)} bugs across {len({p for p, _ in catalog_bugs})} projects.",
        f"Master CSV rows: {len(rows)}; expected matrix rows: {len(catalog_bugs) * len(TECHNIQUES)}.",
        f"Available suite rows: {available_suite_rows}; attempted rows: {attempted_rows}; NO_SUITE rows: {no_suite_rows}; unresolved rows: {unresolved_rows}.",
        "",
        "| Field | Meaning |",
        "|---|---|",
        "| Project, Bug_ID | Defects4J bug key |",
        "| Technique | IPO native, MIO, DeepSeek, or Gemini |",
        "| Target_Classes | Modified classes in scope |",
        "| Suite_Available, Test_Files | Current suite inventory |",
        "| Line_Coverage_%, Branch_Coverage_% | Measured target class coverage; blank unless measured |",
        "| Fault_Detection_Status | BUG_DETECTED only when buggy fails and fixed passes; otherwise measured classification or NOT_EVALUATED |",
        "| Execution_Status | DONE, COMPILE_ERROR, TIMEOUT, NO_SUITE, NOT_RUN, INVALID_SUITE, or STALE_RESULT |",
        "| Suite_SHA256, Run_ID, Timestamp | Suite provenance and run identity |",
        "| Run_Log | Structured JSON run record with provenance, status, timing, and errors or failing-test details |",
        "| Evaluation_Duration_Sec | Wall time of the benchmark evaluation |",
        "",
        "FDR is reported over attempted bug-technique evaluations (DONE, COMPILE_ERROR, TIMEOUT) and separately over the full catalog. A bug is detected only when at least one target test fails on buggy and passes on fixed. Coverage means use DONE rows with measured numeric coverage only.",
        "",
        "MIO budget figures describe suite-generation experiments from the MIO budget summary and are separate from this benchmark evaluation dataset.",
        "",
        "Current execution status by technique:",
        "",
        "| Technique | Execution status | Rows |",
        "|---|---|---:|",
    ]
    for tech in TECHNIQUES:
        for status, count in sorted(status_counts[tech].items()):
            lines.append(f"| {tech} | {status} | {count} |")
    with open(os.path.join(RESULTS, "DATA_DICTIONARY.md"), "w", encoding="utf-8") as stream:
        stream.write("\n".join(lines) + "\n")

    report = [
        "# Member 4 analytics",
        "",
        f"Master dataset currently has {len(rows)} rows for {len(catalog_bugs)} catalog bugs across {len({p for p, _ in catalog_bugs})} projects.",
        ("All available suites have a current benchmark outcome and the full matrix has explicit final statuses."
         if results_complete else
         f"Results are provisional: {attempted_rows}/{available_suite_rows} available suite rows attempted; "
         f"{no_suite_rows} rows have no suite and {unresolved_rows} rows remain unresolved."),
        "",
        "## Bug-level results",
        "",
        "| Technique | Attempted | Completed | Line mean (N) | Branch mean (N) | Detected | FDR of attempts | FDR of full catalog |",
        "|---|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for tech in TECHNIQUES:
        tr = [r for r in rows if r["Technique"] == tech]
        attempts = [r for r in tr if r["Execution_Status"] in ATTEMPTED]
        done = [r for r in tr if r["Execution_Status"] == "DONE"]
        detected = sum(r["Fault_Detection_Status"] == "BUG_DETECTED" for r in attempts)
        line_values = [r["Line_Coverage_%"] for r in done if r["Line_Coverage_%"] is not None]
        branch_values = [r["Branch_Coverage_%"] for r in done if r["Branch_Coverage_%"] is not None]
        fdr = f"{100*detected/len(attempts):.2f}%" if attempts else "n/a"
        full = f"{100*detected/len(catalog_bugs):.2f}%" if catalog_bugs else "n/a"
        line_mean = f"{statistics.mean(line_values):.2f}% (n={len(line_values)})" if line_values else "n/a (n=0)"
        branch_mean = f"{statistics.mean(branch_values):.2f}% (n={len(branch_values)})" if branch_values else "n/a (n=0)"
        report.append(f"| {tech} | {len(attempts)} | {len(done)} | {line_mean} | {branch_mean} | {detected} | {fdr} | {full} |")
    report += ["", "## Coverage", "",
               "Coverage averages include completed evaluations with numeric coverage only; compile failures and missing values are excluded."]
    report += ["", "## Suite generation metrics", "",
               "MIO budget and AI generation records are summarized from their own source files. They are not joined to measured fault detections unless a benchmark run identity supports the join."]
    for tech, values in economics.items():
        report.append(f"- {tech}: {values['generation_records']} generation log records; mean tokens {values['avg_tokens_per_generation_record']}; mean generation time {values['avg_generation_latency_sec']} seconds; measured benchmark detections {values['benchmark_detected_bugs']}.")
    report += ["", "## Paired coverage comparisons", "",
               "Wilcoxon signed-rank tests compare line coverage for the same project-bug keys with measured coverage in both techniques. Holm correction is applied across the six comparisons; the paired rank-biserial correlation reports direction and magnitude."]
    for result in hypothesis:
        if result.get("status") == "MEASURED":
            report.append(f"- {result['group_1']} vs {result['group_2']}: matched n={result['n_pairs']}, nonzero pairs={result['n_nonzero_pairs']}, p={result['p_value']:.4g}, Holm-adjusted p={result['p_value_holm']:.4g}, paired rank-biserial={result['paired_rank_biserial']}.")
        else:
            report.append(f"- {result['group_1']} vs {result['group_2']}: insufficient matched sample (n={result['n_pairs']}).")
    report += ["", "MIO generation budget paired tests compare the same project-bug-target class at two budgets; the output records matched sample size and Holm-adjusted p-values. These tests describe generation coverage and do not measure benchmark fault detection."]
    for comparison, result in budget.get("paired_tests", {}).items():
        if result.get("status") == "MEASURED":
            report.append(f"- MIO {comparison}: matched n={result['n_paired_target_classes']}, p={result['p_value']:.4g}, Holm-adjusted p={result['p_value_holm']:.4g}, paired rank-biserial={result['paired_rank_biserial']}.")
    with open(os.path.join(RESULTS, "advanced_analytics_report.md"), "w", encoding="utf-8") as stream:
        stream.write("\n".join(report) + "\n")


def main():
    rows = load_master()
    catalog_bugs, single_bugs, multi_bugs = load_catalog()
    hypothesis = run_hypothesis(rows)
    budget = run_mio_budget()
    ensemble = run_ensemble(rows, catalog_bugs)
    single_multi = run_single_multi(rows, single_bugs, multi_bugs)
    economics = run_ai_economics(rows)
    export_outputs(rows, catalog_bugs, hypothesis, budget, ensemble, single_multi, economics)
    print(f"Rebuilt analytics from {len(rows)} master rows ({len(catalog_bugs)} catalog bugs).")


if __name__ == "__main__":
    main()
