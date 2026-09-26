from pathlib import Path
import csv, json, statistics, tempfile
from collections import Counter, defaultdict
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TMP = Path(tempfile.gettempdir()) / "member4-report-plan-qa"
OUT = ROOT / "Final_Report.docx"
ROWS = list(csv.DictReader((ROOT / "results/master_benchmark_summary.csv").open(encoding="utf-8-sig", newline="")))
CATALOG = json.loads((ROOT / "target_benchmark/all_bugs_catalog.json").read_text(encoding="utf-8"))
STATS = json.loads((ROOT / "results/master_descriptive_stats.json").read_text(encoding="utf-8"))
ANALYTICS = json.loads((ROOT / "results/advanced_analytics.json").read_text(encoding="utf-8"))
MIO = list(csv.DictReader((ROOT / "MIO_Algorithm/Result_Round2/evosuite_budget_summary.csv").open(encoding="utf-8-sig", newline="")))
AUDIT_ROWS = list(csv.DictReader((ROOT / "results/suite_gap_audit.csv").open(encoding="utf-8-sig", newline="")))
DELIVERY_ROWS = list(csv.DictReader((ROOT / "results/delivery_report_m3.csv").open(encoding="utf-8-sig", newline="")))
TECHS = ["IPO (Native IPO)", "MIO (EvoSuite SBST)", "DeepSeek V4 Flash", "Gemini 3.8 Flash"]
LABEL = {TECHS[0]:"Native IPO", TECHS[1]:"MIO ใน EvoSuite", TECHS[2]:"DeepSeek V4 Flash", TECHS[3]:"Gemini 3.8 Flash"}
CLASS_COUNT = {(x["project"],str(x["bug_id"])):len(x["target_classes"]) for x in CATALOG}
assert len(ROWS)==3416 and len({(r['Project'],r['Bug_ID'],r['Technique']) for r in ROWS})==3416
assert len(CATALOG)==854 and len({r['Project'] for r in ROWS})==17

def value(r,k): return None if r.get(k,'') in ('',None) else float(r[k])
TD={}
for t in TECHS:
    g=[r for r in ROWS if r['Technique']==t]
    attempted=[r for r in g if r['Execution_Status']!='NO_SUITE']
    done=[r for r in g if r['Execution_Status']=='DONE']
    det=[r for r in attempted if r['Fault_Detection_Status']=='BUG_DETECTED']
    line=[value(r,'Line_Coverage_%') for r in done if value(r,'Line_Coverage_%') is not None]
    branch=[value(r,'Branch_Coverage_%') for r in done if value(r,'Branch_Coverage_%') is not None]
    fs=Counter(r['Fault_Detection_Status'] for r in attempted)
    es=Counter(r['Execution_Status'] for r in attempted)
    TD[t]={'attempted':len(attempted),'nosuite':len(g)-len(attempted),'done':es['DONE'],'compile':es['COMPILE_ERROR'],
           'detected':len(det),'notdet':fs['NOT_DETECTED'],'flaky':fs['FLAKY_OR_REGRESSION'],
           'fdr':100*len(det)/len(attempted),'catalog':100*len(det)/854,'n':len(line),
           'lm':statistics.mean(line),'ls':statistics.stdev(line),'bm':statistics.mean(branch),'bs':statistics.stdev(branch)}

EXEC=Counter(r['Execution_Status'] for r in ROWS)
FSTATE=Counter(r['Fault_Detection_Status'] for r in ROWS if r['Execution_Status']!='NO_SUITE')
DELIVERY_KEYS=set()
for item in DELIVERY_ROWS:
    project, bug_text = item['bug'].rsplit('-',1)
    DELIVERY_KEYS.add((project,str(int(bug_text)),item['tech']))
DELIVERY_RESULTS=[r for r in ROWS if (r['Project'],r['Bug_ID'],r['Technique']) in DELIVERY_KEYS]
DELIVERY_EXEC=Counter(r['Execution_Status'] for r in DELIVERY_RESULTS)
DELIVERY_FAULT=Counter(r['Fault_Detection_Status'] for r in DELIVERY_RESULTS)
AI_GAPS=[r for r in AUDIT_ROWS if r['Technique'] in TECHS[2:]]
AI_TARGET_MISMATCHES=[r for r in AI_GAPS if r['Audit_Status']=='TARGET_MISMATCH_CANDIDATE']
AI_MISMATCH_BUGS=', '.join(sorted({f"{r['Project']}-{r['Bug_ID']}" for r in AI_TARGET_MISMATCHES})) or 'ไม่มี'
AI_MISMATCH_TARGETS=', '.join(sorted({r['Target_Classes'] for r in AI_TARGET_MISMATCHES})) or 'ไม่มี'
AI_MISMATCH_PACKAGES=', '.join(sorted({r['Candidate_Packages'] for r in AI_TARGET_MISMATCHES})) or 'ไม่มี'
DETECTED=[r for r in ROWS if r['Fault_Detection_Status']=='BUG_DETECTED']
for r in DETECTED:
    d=json.loads((ROOT/r['Run_Log']).read_text(encoding='utf-8'))
    assert d.get('buggy_failures') and not d.get('fixed_failures')
    assert d.get('suite_sha256')==r['Suite_SHA256'] and d.get('run_id')==r['Run_ID']
CONTR=defaultdict(set)
for r in DETECTED: CONTR[(r['Project'],r['Bug_ID'])].add(r['Technique'])
UNION=set(CONTR)
ATTEMPTED_BUGS={(r['Project'],r['Bug_ID']) for r in ROWS if r['Execution_Status']!='NO_SUITE'}
MIOSTAT={}
for budget in (30,60,120):
    g=[r for r in MIO if int(float(r['Budget_Sec']))==budget]
    cov=[float(r['Line_Cov_Mean_%']) for r in g]
    MIOSTAT[budget]={'n':len(g),'mean':statistics.mean(cov),'sd':statistics.stdev(cov),
                     'seed_sd':statistics.mean(float(r['Line_Cov_SD_%']) for r in g),
                     'duration':statistics.mean(float(r['Avg_Duration_Sec']) for r in g)}

# Create an up-to-date coverage figure from every measured target-class set.
TMP.mkdir(parents=True, exist_ok=True)
W,H=1440,780
chart=Image.new('RGB',(W,H),'white'); draw=ImageDraw.Draw(chart)
font_path=Path(r'C:\Windows\Fonts\arial.ttf')
font=ImageFont.truetype(str(font_path),25); small=ImageFont.truetype(str(font_path),21); titlefont=ImageFont.truetype(str(font_path),30)
left,right,top,bottom=120,1360,115,650
draw.text((left,35),'Mean code coverage — all modified target classes',font=titlefont,fill='#1B3652')
for tick in range(0,101,20):
    y=bottom-(bottom-top)*tick/100
    draw.line((left,y,right,y),fill='#D9E0E6',width=2)
    draw.text((45,y-13),f'{tick}%',font=small,fill='#53606C')
colors=['#2F75B5','#ED7D31']
labels=['Native IPO','MIO','DeepSeek','Gemini']
for j,t in enumerate(TECHS):
    center=left+155+j*305
    for k,key in enumerate(('lm','bm')):
        val=TD[t][key]; x=center-52+k*58
        y=bottom-(bottom-top)*val/100
        draw.rounded_rectangle((x,y,x+42,bottom),radius=6,fill=colors[k])
        draw.text((x-7,y-30),f'{val:.1f}',font=small,fill='#1F2B37')
    name=labels[j]; box=draw.textbbox((0,0),name,font=font)
    draw.text((center-(box[2]-box[0])/2,bottom+18),name,font=font,fill='#1F2B37')
    draw.text((center-45,bottom+55),f"n={TD[t]['n']}",font=small,fill='#53606C')
draw.rectangle((left+390,top-60,left+422,top-30),fill=colors[0]); draw.text((left+432,top-62),'Line',font=small,fill='#1F2B37')
draw.rectangle((left+550,top-60,left+582,top-30),fill=colors[1]); draw.text((left+592,top-62),'Branch',font=small,fill='#1F2B37')
chart.save(TMP/'coverage_all_targets.png')
for src,dst,mask_h in [
    (ROOT/'results/figure2_fdr_distribution.png',TMP/'fault_distribution.png',130),
    (ROOT/'results/figure4_ai_economics.png',TMP/'ai_generation.png',110),
]:
    im=Image.open(src).convert('RGB'); ImageDraw.Draw(im).rectangle((0,0,im.width,mask_h),fill='white'); im.save(dst)

DOC=Document(); SEC=DOC.sections[0]
SEC.page_width,SEC.page_height=Inches(8.5),Inches(11)
SEC.top_margin,SEC.bottom_margin=Inches(.78),Inches(.74)
SEC.left_margin,SEC.right_margin=Inches(.82),Inches(.82)
SEC.different_first_page_header_footer=True
FONT='Tahoma'; NAVY=RGBColor(27,54,82); INK=RGBColor(31,43,55); GRAY=RGBColor(83,96,108)
GRID='D9E0E6'

def sf(st,n,b=False,c=INK):
    st.font.name=FONT; st.font.size=Pt(n); st.font.bold=b; st.font.color.rgb=c
    rp=st.element.get_or_add_rPr(); rf=rp.rFonts
    if rf is None: rf=OxmlElement('w:rFonts'); rp.insert(0,rf)
    for x in ('ascii','hAnsi','eastAsia','cs'): rf.set(qn('w:'+x),FONT)

styles=DOC.styles
sf(styles['Normal'],10.4); styles['Normal'].paragraph_format.line_spacing=1.14; styles['Normal'].paragraph_format.space_after=Pt(5)
sf(styles['Title'],20,True,NAVY); styles['Title'].paragraph_format.space_after=Pt(8)
sf(styles['Subtitle'],11,False,GRAY)
for name,size,bef,aft in [('Heading 1',15,15,7),('Heading 2',12.1,10,5),('Heading 3',10.8,7,3)]:
    st=styles[name]; sf(st,size,True,NAVY); st.paragraph_format.space_before=Pt(bef); st.paragraph_format.space_after=Pt(aft); st.paragraph_format.keep_with_next=True
for name in ('List Bullet','List Number'):
    sf(styles[name],10.3); styles[name].paragraph_format.space_after=Pt(3); styles[name].paragraph_format.line_spacing=1.12
sf(styles['Caption'],9,False,GRAY); styles['Caption'].paragraph_format.space_after=Pt(4); styles['Caption'].paragraph_format.keep_with_next=True
for name in ('TOC 1','TOC 2'):
    try: st=styles[name]
    except KeyError: st=styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    sf(st,8.8,False,INK); st.paragraph_format.space_before=Pt(0); st.paragraph_format.space_after=Pt(0); st.paragraph_format.line_spacing=1.0
tp=styles['Title'].element.get_or_add_pPr()
for el in list(tp):
    if el.tag==qn('w:pBdr'): tp.remove(el)

def field(p,instr,result=''):
    def rchild(tag,attrs=None,text=None):
        r=OxmlElement('w:r'); n=OxmlElement(tag)
        if attrs:
            for k,v in attrs.items(): n.set(k if k.startswith('{') else qn('w:'+k),v)
        if text is not None: n.text=text
        r.append(n); p._p.append(r)
    rchild('w:fldChar',{'fldCharType':'begin'})
    rchild('w:instrText',{'{http://www.w3.org/XML/1998/namespace}space':'preserve'},instr)
    rchild('w:fldChar',{'fldCharType':'separate'})
    if result: rchild('w:t',None,result)
    rchild('w:fldChar',{'fldCharType':'end'})

def heading(text,level=1): return DOC.add_heading(text,level)
def p(text='',style=None,align=None):
    x=DOC.add_paragraph(style=style)
    if text: x.add_run(text)
    if align is not None: x.alignment=align
    return x
def rich(label,text):
    x=DOC.add_paragraph(); x.add_run(label).bold=True; x.add_run(text); return x
def bullet(text): return p(text,'List Bullet')
def number(text): return p(text,'List Number')
def page(): DOC.add_page_break()
def shade(cell,color):
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),color); cell._tc.get_or_add_tcPr().append(sh)
def borders(cell):
    tc=cell._tc.get_or_add_tcPr(); bs=tc.first_child_found_in('w:tcBorders')
    if bs is None: bs=OxmlElement('w:tcBorders'); tc.append(bs)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),GRID); bs.append(e)
def margins(cell):
    tc=cell._tc.get_or_add_tcPr(); m=tc.first_child_found_in('w:tcMar')
    if m is None: m=OxmlElement('w:tcMar'); tc.append(m)
    for side,val in [('top',60),('start',70),('bottom',60),('end',70)]:
        e=OxmlElement('w:'+side); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); m.append(e)
CAPTION_COUNT={'Table':0,'Figure':0}
def caption(kind,text):
    CAPTION_COUNT[kind]+=1
    x=DOC.add_paragraph(style='Caption'); x.add_run('รูปที่ ' if kind=='Figure' else 'ตารางที่ ')
    field(x,' SEQ '+kind+' \\* ARABIC ',str(CAPTION_COUNT[kind])); x.add_run('  '+text)
    x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.keep_with_next=False
    return x
def table(headers,rows,widths,cap,size=8.7):
    t=DOC.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,w in enumerate(widths): t.columns[i].width=Inches(w)
    h=t.rows[0]; h._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for i,txt in enumerate(headers):
        c=h.cells[i]; c.width=Inches(widths[i]); c.text=str(txt); shade(c,'DCE7F0')
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells; t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i,txt in enumerate(row):
            c=cells[i]; c.width=Inches(widths[i]); c.text=str(txt)
            if ridx%2: shade(c,'F6F8FA')
    for ri,row in enumerate(t.rows):
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c); borders(c)
            for q in c.paragraphs:
                q.paragraph_format.space_after=Pt(0); q.paragraph_format.line_spacing=1.08
                for run in q.runs:
                    run.font.name=FONT; run.font.size=Pt(size); run.font.color.rgb=NAVY if ri==0 else INK
                    if ri==0: run.bold=True
    caption('Table',cap)
    DOC.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def figure(path,text,width=6.7):
    q=DOC.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.keep_together=True; q.paragraph_format.keep_with_next=True
    q.add_run().add_picture(str(path),width=Inches(width)); caption('Figure',text)
def ms(a,b): return f'{a:.2f} ± {b:.2f}%'

# Cover
logo_src=ROOT/'assets'/'kku_logo.png'
logo=Image.open(logo_src).convert('RGBA'); logo=logo.crop(logo.getchannel('A').getbbox()); logo_path=TMP/'kku_cover_logo.png'; logo.save(logo_path)
q=DOC.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(8)
q.add_run().add_picture(str(logo_path),height=Inches(1.42))
def cover_line(text,size=12,bold=False,italic=False,before=0,after=3):
    q=DOC.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_before=Pt(before); q.paragraph_format.space_after=Pt(after); q.paragraph_format.line_spacing=1.08
    r=q.add_run(text); r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=INK; r.bold=bold; r.italic=italic
    return q
cover_line('รายงานโครงงานคอมพิวเตอร์',16,True,after=7)
cover_line('การเปรียบเทียบการสร้างชุดทดสอบอัตโนมัติและชุดทดสอบจาก Generative AI บน Defects4J',16,True,after=4)
cover_line('An Empirical Comparison of Automated and Generative AI Test Generation on Defects4J',10.5,False,True,after=12)
cover_line('โดย',12.5,False,before=6,after=5)
cover_line('นายปวริศช์ ประมวล  (รหัสนักศึกษา 673380278-9)',11.5,after=2)
cover_line('นายแทนคุณ พันธ์นิกุล  (รหัสนักศึกษา 673380301-0)',11.5,after=2)
cover_line('นายธนภูมิ จันทรา  (รหัสนักศึกษา 673380272-1)',11.5,after=2)
cover_line('นายศิฆรินทร์ อุปจันทร์  (รหัสนักศึกษา 673380292-5)',11.5,after=10)
cover_line('อาจารย์ประจำวิชา',11.5,False,before=5,after=2)
cover_line('ผู้ช่วยศาสตราจารย์ ดร. ชิตสุธา สุ่มเล็ก',12,after=14)
cover_line('รายงานนี้เป็นส่วนหนึ่งของการศึกษาวิชา CP353201 การประกันคุณภาพซอฟต์แวร์',10.5,after=4)
cover_line('ภาคเรียนที่ 1 ปีการศึกษา 2569',10.5,after=4)
cover_line('สาขาวิทยาการคอมพิวเตอร์ คณะวิทยาศาสตร์',10.5,after=2)
cover_line('มหาวิทยาลัยขอนแก่น',10.5,after=2)
cover_line('(กันยายน พ.ศ. 2569)',10.5,after=0)

# Abstract and TOC
page(); heading('บทคัดย่อ')
p(f"โครงงานนี้เปรียบเทียบ Native IPO, Many-Independent-Objective (MIO) ใน EvoSuite, DeepSeek V4 Flash และ Gemini 3.8 Flash บน Defects4J 3.0.1-7-g8c16da82 ซึ่งมี active bugs 854 บั๊กจาก 17 โครงการ กำหนดผลหนึ่งรายการต่อบั๊กและเทคนิค รวม 3,416 ช่อง มีผลประเมิน suite {len(ROWS)-EXEC['NO_SUITE']:,} ช่อง และ {EXEC['NO_SUITE']:,} ช่องเป็น NO_SUITE. ผล BUG_DETECTED ต้องมีหลักฐานว่า test ล้มเหลวบน buggy version และผ่านบน fixed version.")
p(f"Coverage รวมทุก modified target class ตามค่า aggregate ใน Defects4J summary. Gemini มี line/branch coverage เฉลี่ยสูงสุด {TD[TECHS[3]]['lm']:.2f}%/{TD[TECHS[3]]['bm']:.2f}% (n={TD[TECHS[3]]['n']}). Native IPO ตรวจพบ {TD[TECHS[0]]['detected']}/{TD[TECHS[0]]['attempted']} suite ({TD[TECHS[0]]['fdr']:.2f}%), MIO {TD[TECHS[1]]['detected']}/{TD[TECHS[1]]['attempted']} ({TD[TECHS[1]]['fdr']:.2f}%), DeepSeek {TD[TECHS[2]]['detected']}/{TD[TECHS[2]]['attempted']} ({TD[TECHS[2]]['fdr']:.2f}%) และ Gemini {TD[TECHS[3]]['detected']}/{TD[TECHS[3]]['attempted']} ({TD[TECHS[3]]['fdr']:.2f}%).")
p(f"ผลรวมสี่เทคนิคตรวจพบ {len(UNION)} บั๊กไม่ซ้ำจาก {len(ATTEMPTED_BUGS)} บั๊กที่มี suite อย่างน้อยหนึ่งเทคนิค ({100*len(UNION)/len(ATTEMPTED_BUGS):.2f}%). สถิติ MIO budget และ token/เวลา AI แยกจาก benchmark กลางเพราะ generation logs ไม่มี run ID เชื่อมกับผลตรวจจับรายบั๊ก จึงไม่คำนวณ token ต่อบั๊กที่ตรวจพบ.")
rich('คำสำคัญ  ','Defects4J, software testing, IPO, MIO, EvoSuite, Generative AI, code coverage, fault detection')
page(); heading('สารบัญ'); q=DOC.add_paragraph(); field(q,' TOC \\o "1-1" \\h \\z \\u ','สารบัญอัตโนมัติ — อัปเดตใน Word')
page(); heading('สารบัญตาราง'); q=DOC.add_paragraph(); field(q,' TOC \\h \\z \\c "Table" ','รายการตารางอัตโนมัติ — อัปเดตใน Word')
heading('สารบัญภาพ'); q=DOC.add_paragraph(); field(q,' TOC \\h \\z \\c "Figure" ','รายการภาพอัตโนมัติ — อัปเดตใน Word')

# Chapter 1
page(); heading('บทที่ 1 บทนำและขอบเขตการศึกษา')
heading('1.1 ที่มาและความสำคัญ',2)
p('การสร้าง unit test ทำได้หลายแนวทาง ตั้งแต่การออกแบบค่าป้อนเข้าแบบ combinatorial การค้นหาด้วย fitness function ไปจนถึงการใช้ Generative AI สร้างโค้ดทดสอบ. จำนวน test หรือ coverage สูงเพียงอย่างเดียวไม่ยืนยันว่า test ตรวจจับข้อบกพร่องจริง งานนี้จึงประเมิน suite จากทั้งสี่แนวทางบนโปรแกรม Java ที่มีบั๊กจริง โดยแยก code coverage ออกจากการทดสอบว่า suite แยก buggy และ fixed version ได้หรือไม่.')
p('Defects4J จัดเตรียมบั๊กจริงและเครื่องมือ checkout เวอร์ชัน buggy/fixed เพื่อสนับสนุนการศึกษาซอฟต์แวร์เชิงทดลอง (Just et al., 2014). งานนี้ใช้ Defects4J รุ่น 3.0.1-7-g8c16da82 ซึ่งมี 854 active bugs จาก 17 โครงการ (Defects4J contributors, n.d.) และประเมินเฉพาะ suite ที่มีหลักฐานและผ่านเกณฑ์คัดเลือกของ runner.')
heading('1.2 วัตถุประสงค์',2)
for x in ['เปรียบเทียบ Native IPO, MIO ใน EvoSuite, DeepSeek V4 Flash และ Gemini 3.8 Flash ด้วย runner กลางชุดเดียวกัน',
          'วัด line และ branch coverage ในขอบเขต modified target classes พร้อมแสดงจำนวนข้อมูลที่ใช้คำนวณ',
          'วัด bug-level fault detection โดยตรวจผลบน buggy/fixed และแยก compile error กับ flaky/regression',
          'วิเคราะห์ผลรวม ensemble และข้อมูลสร้าง suite เช่น budget, seed, token และเวลา โดยรักษาที่มาของข้อมูล',
          'จัดเก็บ source, test suite, prompt, configuration, logs และคำสั่งทำซ้ำใน repository']: bullet(x)
heading('1.3 ขอบเขต',2)
p('Catalog มี 854 บั๊ก, 17 โครงการ, 1,073 modified target class records และ 577 คลาสไม่ซ้ำ. มี 727 บั๊กที่แก้คลาสเดียว และ 127 บั๊กที่แก้หลายคลาส. โครงการคือ Chart, Cli, Closure, Codec, Collections, Compress, Csv, Gson, JacksonCore, JacksonDatabind, JacksonXml, Jsoup, JxPath, Lang, Math, Mockito และ Time.')
p('หน่วยวิเคราะห์หลักคือ project + bug_id + technique; หนึ่งบั๊กนับหนึ่งครั้งต่อเทคนิค แม้มีหลายคลาสเป้าหมาย. Coverage ใช้ aggregate counts ของทุก modified target class ที่ Defects4J เขียนใน summary.csv จึงคำนวณจากผลรวม covered ÷ total ของบรรทัดและ condition branches ไม่ใช่ค่าเฉลี่ยรายคลาส. จำนวน coverage n นับเฉพาะผล DONE ที่มีค่าตัวเลขจริง.')
heading('1.4 คำถามวิจัย',2)
table(['คำถาม','วิธีตอบ'],[
    ('RQ1 Coverage','เทียบ line/branch mean ของทุก modified target class ที่วัดได้ พร้อมระบุ n'),
    ('RQ2 Fault detection','เทียบ BUG_DETECTED และ FDR ต่อ suite ที่รัน โดยรวม compile/flaky ในตัวหาร'),
    ('RQ3 Ensemble','นับบั๊กไม่ซ้ำที่อย่างน้อยหนึ่งเทคนิคตรวจพบ และส่วนร่วมเฉพาะของแต่ละเทคนิค'),
    ('RQ4 MIO budget','สรุป coverage criterion รวมและเวลาแยกตาม budget พร้อม seed และข้อจำกัด'),
    ('RQ5 AI generation','เปรียบเทียบ token/เวลาเฉลี่ยต่อ generation record แยกจาก benchmark detection')
],[1.45,5.25],'คำถามวิจัยและหลักฐานที่ใช้ตอบ',8.8)

# Chapter 2
page(); heading('บทที่ 2 แนวทางสร้างชุดทดสอบ')
heading('2.1 Native IPO และ pairwise testing',2)
p('Pairwise testing ต้องการให้ทุกคู่ค่าของ input factors ปรากฏอย่างน้อยหนึ่งครั้ง. IPO (In-Parameter-Order) สร้าง covering array แบบเพิ่มพารามิเตอร์ทีละตัว โดยขยายแถวเดิมในแนวนอนและเพิ่มแถวในแนวตั้งเมื่อยังมีคู่ค่าที่ขาด (Lei & Tai, 1998). Pairwise ช่วยควบคุมจำนวนกรณี แต่ไม่รับประกันการตรวจพบบั๊ก จึงวัด fault detection แยกจาก pair coverage.')
p('สาย IPO ใช้ Native IPO pipeline ใน `Combinatorial_IPO/Code/` เพื่อวิเคราะห์ Java API, เตรียม construction/value domains, สร้าง 2-way combinations, สร้าง oracle จาก fixed version และปล่อย JUnit 4 suite เมื่อผ่าน verification. Manifest `Combinatorial_IPO/Results/verified_suites_manifest.json` เก็บ suite hash และสถานะตรวจ. PICT เป็นข้อมูลอ้างอิงของสาย IPO ไม่ใช่ผล IPO หลักใน benchmark นี้; ตัวเลข 173 suites และ 42,398 tests เป็น baseline คนละรอบ.')
heading('2.2 MIO และ EvoSuite',2)
p('MIO เป็น search algorithm สำหรับการสร้าง test suite ที่จัดการเป้าหมาย coverage หลายรายการแยกอิสระและปรับการค้นหาตาม budget (Arcuri, 2018). EvoSuite สร้าง unit tests สำหรับ Java ด้วย search-based approaches และ assertions (Fraser & Arcuri, 2011).')
p('Budget experiment ใช้ EvoSuite 1.0.6, algorithm MIO, criterion `LINE:BRANCH`, budgets 30/60/120 วินาที และ seeds 101, 102, 103. ผล generation แยกจากผล benchmark กลาง. Raw criterion ให้ coverage ค่าเดียว จึงไม่รายงานเป็น line และ branch แยกกัน.')
heading('2.3 DeepSeek และ Gemini',2)
p('สคริปต์ `scripts/kku_generate.py` เรียก KKU IntelSphere API โดยใช้ default identifiers `deepseek-v4-flash` และ `gemini-3.8-flash`. ชื่อโมเดลในรายงานเป็น labels ที่บันทึกในโครงงาน; generation records ไม่เก็บ provider model ID ทุกแถว จึงไม่อนุมานสถาปัตยกรรมเฉพาะรุ่น.')
p('Prompt ใช้ source code ของ target class และ defect context ที่มีใน repository. Prompt กำหนด JUnit 4, package ให้ตรงกับโครงการ, test timeout, หลีกเลี่ยง dependency mocking ภายนอก และชี้นำให้ทดสอบค่าขอบเขต. Prompt ที่ใช้จริงเก็บราย bug ใน `Deepseek-v4_flash/Prompt/` และ `Gemini-3_8_flash/Prompt/`.')
table(['แนวทาง','เครื่องมือ/วิธี','หลักฐานการสร้าง'],[
    ('Combinatorial','Native IPO, 2-way pair coverage','Generator source, verified manifest, suite hash'),
    ('Search-based','EvoSuite 1.0.6, MIO','Budget/seed summary และ raw reports'),
    ('Generative AI','DeepSeek V4 Flash label','Prompt, JUnit test, generation records'),
    ('Generative AI','Gemini 3.8 Flash label','Prompt, JUnit test, generation records')
],[1.25,2.3,3.15],'วิธีสร้าง suite ที่นำมาประเมิน',8.7)

# Chapter 3
page(); heading('บทที่ 3 สภาพแวดล้อมและวิธีทดลอง')
heading('3.1 แหล่งข้อมูลและการจัดการผล',2)
p(f"Catalog อยู่ใน `target_benchmark/all_bugs_catalog.json`; master results อยู่ใน `results/master_benchmark_summary.csv`; descriptive statistics อยู่ใน `results/master_descriptive_stats.json`. Master มี 3,416 คีย์ไม่ซ้ำ ครบ 854 บั๊ก × 4 เทคนิค. มี {len(ROWS)-EXEC['NO_SUITE']:,} คีย์ที่ runner ประเมินแล้ว; {EXEC['NO_SUITE']:,} คีย์เป็น NO_SUITE; unresolved outcomes {sum(EXEC[s] for s in ('NOT_RUN','STALE_RESULT','CHECKOUT_ERROR','INVALID_SUITE','RUN_ERROR'))}. บัญชีตรวจเพิ่มเติมใน `results/suite_gap_audit.csv` แยกกรณีไม่มีไฟล์ Javaออกจากไฟล์ผู้สมัครที่ยังไม่ผ่าน manifest หรือ target matcher.")
p('Runner ตรวจ target classes กับ Defects4J, เลือก suite ตาม bug key และเทคนิค และเก็บ SHA-256, run ID, timestamp, duration และ structured log. ผลมาตรฐานหนึ่งรายการต่อคีย์ถูกรวมใน master; run ซ้ำไม่ได้เพิ่มจำนวนบั๊กในตัวหาร.')
heading('3.2 สภาพแวดล้อม',2)
table(['องค์ประกอบ','การตั้งค่าที่อ้างอิง'],[
    ('Container','Docker service `defects4j_sqa`; timezone Asia/Bangkok'),
    ('Benchmark','Defects4J 3.0.1-7-g8c16da82; 854 active bugs ตามรุ่นที่ติดตั้ง'),
    ('Java','ติดตั้ง OpenJDK 8 และ 11; default JAVA_HOME เป็น JDK 11; MIO batch ตั้ง JDK 8'),
    ('Test framework','JUnit 4; test prompt กำหนด timeout 4 วินาที'),
    ('Coverage','Defects4J coverage workflow พร้อม Cobertura summary')
],[1.45,5.25],'สภาพแวดล้อมการประเมิน',8.8)
p('Dockerfile และ compose อยู่ใน `docker/Dockerfile` และ `docker/docker-compose.yml`. Defects4J CLI ใช้ JDK 11; MIO generation บางขั้นตอนตั้ง JDK 8. Dockerfile สร้างจาก Ubuntu 20.04 พร้อมตรึง commit ของ Defects4J และ PICT จึงไม่ต้องเตรียม image พื้นฐานเฉพาะของเครื่องก่อน build. การ build ครั้งแรกต้องดาวน์โหลดและ initialize catalog ของ Defects4J ซึ่งใช้พื้นที่และเวลามาก; แพ็กเกจระบบยังติดตั้งจาก Ubuntu package archive.')
heading('3.3 ขั้นตอนประเมิน',2)
for x in ['อ่าน target classes จาก catalog และตรวจเทียบกับ Defects4J',
          'ค้นหา suite ตาม project, bug และ technique; หากไม่มีให้บันทึก NO_SUITE',
          'คอมไพล์ suite และวัด coverage ผ่าน Defects4J workflow บน buggy version',
          'รัน suite บน buggy; ตรวจ fixed version ด้วย suite เดียวกันเพื่อยืนยัน defect-specific failure',
          'บันทึก status, filename, suite hash, run ID, timestamp, duration และ log',
          'รวมผลเป็น bug-technique matrix แล้วคำนวณสถิติจากคีย์ไม่ซ้ำ']: number(x)
heading('3.4 นิยามสถานะและตัวชี้วัด',2)
table(['สถานะ','นิยามที่ใช้'],[
    ('BUG_DETECTED','test อย่างน้อยหนึ่งกรณี fail บน buggy และไม่มี failure บน fixed; ตรวจจาก log'),
    ('NOT_DETECTED','suite รันได้แต่ไม่ทำให้ buggy fail'),
    ('FLAKY_OR_REGRESSION','failure ยังเกิดบน fixed หรือไม่จำเพาะต่อบั๊ก'),
    ('COMPILE_ERROR','suite คอมไพล์ไม่ผ่าน; ไม่ถือเป็น coverage 0 ที่วัดได้'),
    ('NO_SUITE','ไม่มี suite ให้ประเมิน; ไม่ถือเป็น NOT_DETECTED')
],[1.75,4.95],'นิยามสถานะผลการประเมิน',8.8)
p('FDR ต่อ suite = BUG_DETECTED ÷ suite ที่นำมาประเมิน; compile errors และ flaky/regression อยู่ในตัวหาร. อัตรา detected ÷ 854 แสดงเพิ่มเติมเพื่อสื่อสัดส่วนการตรวจพบใน catalog ทั้งหมด; NO_SUITE ไม่ถูกเรียกว่าเป็นผลทดสอบ.')

# Chapter 4
heading('บทที่ 4 ผลการทดลอง')
heading('4.1 ความครบถ้วนและสถานะผล',2)
p(f"Catalog มี 854 บั๊ก × 4 เทคนิค รวม 3,416 คีย์ไม่ซ้ำ. มี suite และผลประเมิน {len(ROWS)-EXEC['NO_SUITE']:,} รายการ; {EXEC['NO_SUITE']:,} รายการไม่มี suite. สถานะมี {EXEC['DONE']:,} DONE, {EXEC['COMPILE_ERROR']:,} COMPILE_ERROR และ {EXEC['TIMEOUT']:,} TIMEOUT. ในผลที่รันได้มี {FSTATE['BUG_DETECTED']:,} BUG_DETECTED, {FSTATE['NOT_DETECTED']:,} NOT_DETECTED และ {FSTATE['FLAKY_OR_REGRESSION']:,} FLAKY_OR_REGRESSION; unresolved outcomes {sum(EXEC[s] for s in ('NOT_RUN','STALE_RESULT','CHECKOUT_ERROR','INVALID_SUITE','RUN_ERROR'))}.")
table(['เทคนิค','มี suite','NO_SUITE','DONE','compile','พบ defect','flaky','ไม่พบ','FDR suite','พบ/854'],[
    (LABEL[t],f"{TD[t]['attempted']}/854",TD[t]['nosuite'],TD[t]['done'],TD[t]['compile'],TD[t]['detected'],TD[t]['flaky'],TD[t]['notdet'],f"{TD[t]['fdr']:.2f}%",f"{TD[t]['catalog']:.2f}%") for t in TECHS
],[1.16,.58,.62,.42,.61,.59,.46,.44,.66,.54],'สถานะและ fault detection ต่อเทคนิค; FDR ต่อ suite รวม compile/flaky ในตัวหาร',7.0)
p(f"`DONE` หมายถึงการประเมินสิ้นสุดโดยไม่เกิด compile error; fault status ของแถว DONE ยังอาจเป็น flaky/regression. จำนวน {EXEC['NO_SUITE']:,} NO_SUITE แสดงแยกจาก NOT_DETECTED.")
figure(TMP/'fault_distribution.png','สัดส่วนสถานะใน suite ที่นำมาประเมิน; n ระบุใต้แต่ละเทคนิค')
heading('4.2 Coverage ของ modified target classes',2)
p('Defects4J coverage workflow เขียนจำนวนบรรทัดและ condition branches ที่ครอบคลุมและทั้งหมดเป็น aggregate summary ของคลาสที่ถูกวัด. รายงานจึงรวมบั๊กที่แก้หนึ่งคลาสและหลายคลาส โดยใช้ค่าร้อยละที่คำนวณจาก aggregate counts. ค่า n ต่างกันเพราะ suite ที่ไม่มีหรือ compile ไม่ผ่านไม่มี coverage ที่วัดได้. SD แสดงการกระจายระหว่างบั๊ก ไม่ใช่ standard error.')
covrows=[(LABEL[t],TD[t]['n'],ms(TD[t]['lm'],TD[t]['ls']),ms(TD[t]['bm'],TD[t]['bs'])) for t in TECHS]
table(['เทคนิค','n','Line mean ± SD','Branch mean ± SD'],covrows,[1.7,.55,2.2,2.25],'Coverage ของ modified target classes ที่มีค่าการวัด',8.6)
figure(TMP/'coverage_all_targets.png','Line และ branch coverage เฉลี่ยของทุก modified target class ที่วัดได้; n คือจำนวนบั๊ก')
p(f"Gemini มี coverage mean สูงสุดในผลที่วัดได้. ค่าเฉลี่ยนี้ยังมี selection bias ได้ เพราะ coverage คำนวณได้เฉพาะกรณีที่ suite ผ่านการ compile และทำงานสำเร็จ; Gemini compile error {TD[TECHS[3]]['compile']}/{TD[TECHS[3]]['attempted']} และ DeepSeek {TD[TECHS[2]]['compile']}/{TD[TECHS[2]]['attempted']}.")
heading('4.3 การเปรียบเทียบ line coverage แบบจับคู่',2)
PAIR_LABEL = {name: LABEL[name] for name in TECHS}
paired = ANALYTICS['hypothesis_testing']
paired_rows=[]
for item in paired:
    left_name=PAIR_LABEL.get(item['group_1'],item['group_1'])
    right_name=PAIR_LABEL.get(item['group_2'],item['group_2'])
    p_holm=item.get('p_value_holm')
    paired_rows.append((f"{left_name} – {right_name}",item.get('n_pairs','—'),
                        f"{p_holm:.2e}" if p_holm is not None else '—',
                        f"{item['paired_rank_biserial']:.3f}" if item.get('paired_rank_biserial') is not None else '—'))
table(['คู่เปรียบเทียบ','บั๊กที่มีคู่ข้อมูล','p หลัง Holm','Rank-biserial'],paired_rows,[3.0,1.05,1.25,1.4],'ผล Wilcoxon signed-rank สำหรับ line coverage ที่วัดได้',8.0)
p('เปรียบเทียบเฉพาะ project–bug ที่ทั้งสองเทคนิคมี line coverage และปรับ p-value ด้วย Holm สำหรับหกคู่. ค่า rank-biserial เป็นบวกเมื่อเทคนิคซ้ายมี coverage สูงกว่า; ผลนี้เป็นการเปรียบเทียบเชิงสำรวจของ suite ที่ผ่านการวัดและไม่ลบข้อจำกัดจากจำนวน compile error ที่ต่างกัน.')
heading('4.4 Ensemble',2)
p(f"Union ของ BUG_DETECTED มี {len(UNION)} บั๊กไม่ซ้ำจาก {len(ATTEMPTED_BUGS)} บั๊กที่มี suite อย่างน้อยหนึ่งเทคนิค ({100*len(UNION)/len(ATTEMPTED_BUGS):.2f}%) และคิดเป็น {100*len(UNION)/854:.2f}% ของ 854 บั๊ก. นี่เป็นการรวมผลที่วัด ไม่ใช่การรัน ensemble suite ใหม่.")
uc=[]
for t in TECHS:
    unique=sum(1 for _,ts in CONTR.items() if t in ts and len(ts)==1)
    uc.append((LABEL[t],TD[t]['detected'],unique))
table(['เทคนิค','พบทั้งหมด','พบเฉพาะเทคนิค'],uc,[2.25,1.35,3.0],'จำนวนการตรวจพบทั้งหมดและการมีส่วนร่วมที่ไม่ซ้ำ',8.8)

# Chapter 5
page(); heading('บทที่ 5 การวิเคราะห์และอภิปรายผล')
heading('5.1 Coverage และ fault detection',2)
p('Gemini มีค่า coverage สูงสุดในชุดผลที่วัดได้ ขณะที่ Native IPO มี FDR ต่อ suite สูงสุด. MIO และ DeepSeek แสดงว่าการเข้าถึงโค้ดไม่ได้รับประกันว่า assertion จะจับพฤติกรรมบกพร่องได้. Coverage เป็นตัวชี้วัดการเข้าถึงโค้ด; fault detection วัดว่า suite แยก buggy จาก fixed ได้หรือไม่.')
p(f"Native IPO ตรวจพบ {TD[TECHS[0]]['detected']}/{TD[TECHS[0]]['attempted']} suite ({TD[TECHS[0]]['fdr']:.2f}%); ต่อ catalog เท่ากับ {TD[TECHS[0]]['detected']}/854 ({TD[TECHS[0]]['catalog']:.2f}%). มี suite ให้ประเมิน {TD[TECHS[0]]['attempted']} บั๊ก. Gemini ตรวจพบ {TD[TECHS[3]]['detected']}/{TD[TECHS[3]]['attempted']} ({TD[TECHS[3]]['fdr']:.2f}%) และมี compile error {TD[TECHS[3]]['compile']} รายการ. การรายงานทั้งตัวหาร suite และตัวหาร catalog ป้องกันการอ่านเปอร์เซ็นต์โดยไม่เห็นขอบเขต.")
heading('5.2 Compile errors และความไม่เสถียร',2)
p(f"ใน suite evaluations {EXEC['COMPILE_ERROR']:,} รายการ compile ไม่ผ่าน และ {FSTATE['FLAKY_OR_REGRESSION']:,} รายการเป็น flaky/regression. DeepSeek มี compile errors {TD[TECHS[2]]['compile']}/{TD[TECHS[2]]['attempted']} และ Gemini {TD[TECHS[3]]['compile']}/{TD[TECHS[3]]['attempted']}. สถานะเหล่านี้ไม่ถูกแทนด้วย coverage 0; FDR ต่อ suite รวมไว้ในตัวหารเพื่อให้เห็นความพร้อมใช้จริง.")
heading('5.3 MIO search budget',2)
p('MIO budget summary เก็บผลตาม target class, budget และสาม seeds. Raw EvoSuite criterion `LINE;BRANCH` ให้ coverage ค่าเดียว; สคริปต์บันทึกค่าเดียวกันในคอลัมน์ line และ branch ดังนั้นรายงานเป็น combined coverage และไม่แสดง branch แยกหรือ mutation score ซึ่งใน summary เป็นค่าคงที่ placeholder. n คือจำนวน class-budget records ไม่ใช่จำนวนบั๊กไม่ซ้ำ.')
br=[]
for b in (30,60,120):
    d=MIOSTAT[b]
    delta='ฐาน' if b==30 else f"+{d['mean']-MIOSTAT[b//2]['mean']:.2f} จุด" if b==60 else f"+{d['mean']-MIOSTAT[60]['mean']:.2f} จุด"
    br.append((f'{b} วินาที',d['n'],ms(d['mean'],d['sd']),f"{d['seed_sd']:.2f} จุด",f"{d['duration']:.2f} วินาที",delta))
table(['Budget','records','Coverage รวม mean ± SD','SD ใน 3 seeds เฉลี่ย','เวลาเฉลี่ย','เทียบก่อน'],br,[.9,.68,1.7,1.42,1.17,1.25],'ข้อมูลการสร้าง MIO suite แยกตาม budget',7.9)
p('SD coverage ในตารางแสดงการกระจายระหว่าง class-budget records; SD ใน seed column คือค่าเฉลี่ยของความผันผวนภายใน record จากสาม seeds. เวลาเฉลี่ยรวม setup/compile จึงยาวกว่า search budget ได้. การทดสอบ Wilcoxon สำหรับ MIO จับคู่เฉพาะ project–bug–target class ที่มีผลทั้งสอง budget และปรับ p-value ด้วย Holm. ผล budget เป็น generation experiment แยกจาก benchmark กลางและใช้เป็นคำอธิบายเชิงพรรณนา ไม่ใช่ FDR.')
mio_pairs=ANALYTICS['mio_generation_budget']['paired_tests']
mio_30_60=mio_pairs['30s_vs_60s']; mio_60_120=mio_pairs['60s_vs_120s']
p(f"เมื่อจับคู่ target classes ที่มีผลทั้งสอง budget พบความต่างของ coverage ระหว่าง 30 กับ 60 วินาที (n={mio_30_60['n_paired_target_classes']}, p หลัง Holm={mio_30_60['p_value_holm']:.2e}, rank-biserial={mio_30_60['paired_rank_biserial']:.3f}) และระหว่าง 60 กับ 120 วินาที (n={mio_60_120['n_paired_target_classes']}, p หลัง Holm={mio_60_120['p_value_holm']:.2e}, rank-biserial={mio_60_120['paired_rank_biserial']:.3f}). rank-biserial ติดลบเพราะคำนวณผลต่าง budget สั้นลบ budget ยาว; ค่านี้จึงชี้ว่า coverage ใน budget ยาวสูงกว่าในคู่ที่วัดได้.")
heading('5.4 AI token และเวลา generation',2)
ai=ANALYTICS['ai_economics']
air=[('Gemini 3.8 Flash',ai['Gemini 3.8 Flash']['generation_records'],f"{ai['Gemini 3.8 Flash']['avg_tokens_per_generation_record']:,.2f}",f"{ai['Gemini 3.8 Flash']['avg_generation_latency_sec']:.2f} วินาที"),
     ('DeepSeek V4 Flash',ai['DeepSeek V4 Flash']['generation_records'],f"{ai['DeepSeek V4 Flash']['avg_tokens_per_generation_record']:,.2f}",f"{ai['DeepSeek V4 Flash']['avg_generation_latency_sec']:.2f} วินาที")]
table(['โมเดลตาม label','records','Token/record เฉลี่ย','เวลา/record เฉลี่ย'],air,[1.85,1.15,1.85,2.0],'สถิติจาก log การสร้างชุดทดสอบ AI',8.7)
figure(TMP/'ai_generation.png','Token consumption และ generation latency เฉลี่ยต่อ record ตาม log')
p('Generation record ไม่ใช่ benchmark evaluation และไม่มี run ID ผูกกับผลตรวจพบรายบั๊ก จึงห้ามใช้หารเป็น token ต่อการตรวจพบ. Provider model ID ไม่ได้เก็บครบทุกแถว จึงรายงานชื่อรุ่นตาม label ในโครงงาน.')
heading('5.5 ข้อจำกัด',2)
for x in [f"จำนวน suite ต่างกันระหว่างเทคนิค; {EXEC['NO_SUITE']:,} ช่อง NO_SUITE เป็นช่องที่ไม่มี suite สำหรับ benchmark evaluation และไม่ใช่ผลลบของการทดสอบ. IPO มี {TD[TECHS[0]]['nosuite']} ช่อง: 37 GENERATION_OR_VERIFICATION_ERROR และ 560 SKIPPED_NOT_READY. ใน 37 ช่องมี candidate Java 6 ช่องที่ fixed-version verification ไม่ผ่าน และอีก 31 ช่องไม่พบ candidate; ใน 560 ช่องที่ถูกข้าม มี 495 NEEDS_ADAPTER, 47 NEEDS_ENTRY_POINT, 16 ที่มีทั้งสองสาเหตุ และ 2 ที่มี NEEDS_ADAPTER ร่วมกับ NOT_PAIRWISE_APPLICABLE.",
            f"MIO มี {TD[TECHS[1]]['nosuite']} GENERATION_FAILURE หลังบันทึกการลองสร้าง 3 budgets คูณ 3 seeds ต่อบั๊ก: Mockito 15 ช่อง compile ไม่ผ่านจาก Bintray/JCenter; Math-13, Math-31 และ Gson-3 เกิด EvoSuite 1.0.6 internal NPE; Gson-8 เกิด JVM native crash; JacksonDatabind-24 เกิดปัญหา character encoding ตอน compile. ทุกช่องยังคงเป็น NO_SUITE ใน benchmark เพราะไม่มีไฟล์ suite สำหรับประเมิน.",
            f"Member 3 ส่งไฟล์ AI {len(DELIVERY_ROWS)} ไฟล์สำหรับ {len(DELIVERY_KEYS)} คีย์บั๊ก–เทคนิค; runner ประเมินครบแล้ว ได้ COMPILE_ERROR {DELIVERY_EXEC['COMPILE_ERROR']} คีย์ และ DONE {DELIVERY_EXEC['DONE']} คีย์ โดย {DELIVERY_FAULT['FLAKY_OR_REGRESSION']} คีย์ในกลุ่ม DONE ยัง fail บน fixed. เหลือ AI NO_SUITE {len(AI_GAPS)} คีย์ที่ target mismatch ({AI_MISMATCH_BUGS}); catalog/Defects4J ระบุ target {AI_MISMATCH_TARGETS} แต่ candidate ใช้ package {AI_MISMATCH_PACKAGES}.",
            'Coverage aggregate รวม modified target classes จาก Defects4J summary; การเฉลี่ยนี้ยังขึ้นกับชุด suite ที่ compile และรันสำเร็จ.',
          'Compile errors และ flaky/regression สูงใน AI suites; coverage mean จากชุดที่วัดได้อาจมี survivorship bias.',
          'MIO budget summary ใช้ criterion coverage รวม ไม่แยก line/branch; cohort ต่อ budget ไม่เหมือนกัน.',
          'AI generation log ไม่มี run ID ครบ; ไม่รองรับการวิเคราะห์ token/time ต่อผล detection.',
          'Benchmark evaluation ต่อ bug-technique key ไม่ได้ทำซ้ำทุก suite หลายรอบเพื่อวัด variance.',
          'ข้อสรุปจำกัดอยู่ที่ Java Defects4J catalog และ suites ที่ทีมส่งมอบใน snapshot นี้.']: bullet(x)

# Chapter 6
page(); heading('บทที่ 6 สรุปผลและการทำซ้ำ')
heading('6.1 สรุปคำตอบ',2)
for x in ['RQ1: Gemini มี line/branch coverage mean สูงสุดในผลที่วัดได้เมื่อรวม modified target classes ทุกคลาส; n แสดงใต้ตารางและผล compile error แสดงแยก.',
          f"RQ2: Native IPO มี FDR ต่อ suite สูงสุดที่ {TD[TECHS[0]]['fdr']:.2f}%; Gemini ตรวจพบจำนวนบั๊กมากที่สุด {TD[TECHS[3]]['detected']}/{TD[TECHS[3]]['attempted']} ({TD[TECHS[3]]['fdr']:.2f}%). ตัวหารต่างกัน.",
          f"RQ3: Union พบ {len(UNION)} บั๊กไม่ซ้ำ; ส่วนร่วมเฉพาะคือ Gemini {sum(1 for _,ts in CONTR.items() if TECHS[3] in ts and len(ts)==1)}, IPO {sum(1 for _,ts in CONTR.items() if TECHS[0] in ts and len(ts)==1)}, MIO {sum(1 for _,ts in CONTR.items() if TECHS[1] in ts and len(ts)==1)} และ DeepSeek {sum(1 for _,ts in CONTR.items() if TECHS[2] in ts and len(ts)==1)}.",
          'RQ4: Coverage criterion รวมใน MIO summary เพิ่มตาม budget เชิงพรรณนา; cohort ต่างกันจึงไม่ยืนยัน causal gain หรือจุดอิ่มตัว.',
          'RQ5: AI logs ให้ token และเวลาเฉลี่ยต่อ generation record; ไม่มี run ID เชื่อมกับ detection จึงไม่คำนวณต้นทุนต่อบั๊ก.']: bullet(x)
heading('6.2 ข้อสรุป',2)
p(f"ไม่มีวิธีเดียวดีที่สุดทุกตัวชี้วัด. Gemini ให้ coverage สูงสุดในผลที่วัดได้และตรวจพบ {TD[TECHS[3]]['detected']} บั๊ก แต่มี compile errors จำนวนมาก. Native IPO มี FDR ต่อ suite สูงสุดใน snapshot นี้ แต่มี suite ผ่านเกณฑ์ประเมินน้อยกว่า. ผล MIO และ DeepSeek แสดงว่า coverage สูงไม่ได้รับประกัน fault detection. การเลือกใช้จึงควรพิจารณา coverage, detection, ความสามารถ compile และ stability ร่วมกัน.")
heading('6.3 งานต่อไป',2)
for x in ['Member 3 ตรวจ Math-13: แก้ package ให้ตรงกับ classes.modified ที่ Defects4J ยืนยัน หรือยืนยันว่าไม่มี suite ใช้ได้เพื่อคงสถานะ NO_SUITE; หากส่งไฟล์ใหม่ Member 4 ประเมินเฉพาะคีย์ที่เปลี่ยน.',
          'บันทึก immutable model ID และ generation record ID เชื่อมกับ suite hash และ benchmark run ID.',
          'แยก MIO raw reports ตาม class, budget และ seed ให้มี provenance ครบทุกแถว; แยก coverage criterion ที่วัดได้จริง.',
          'ปรับ prompt และเพิ่ม compile feedback เพื่อแก้ compile errors และ flaky assertions.',
          'ทำ repeated benchmark executions หากต้องการรายงาน variance ของผลต่อ suite.']: bullet(x)
heading('6.4 คำสั่งทำซ้ำ',2)
p('คำสั่งต่อไปนี้อ้างอิง README และ Docker configuration. การประเมินครบทุก suite ใช้เวลาตามจำนวน suite และ environment; หากต้องการสร้าง analytics จากผลที่บันทึกไว้สามารถข้าม runner ได้.')
for cmd in ['docker compose -f docker/docker-compose.yml up -d --build',
            'docker exec defects4j_sqa python3 /workspace/scripts/run_benchmark.py --all-bugs --resume',
            'python scripts/consolidate_master_results.py',
            'python scripts/advanced_data_analytics.py']:
    q=p(); q.paragraph_format.left_indent=Inches(.25); q.paragraph_format.space_after=Pt(3)
    r=q.add_run(cmd); r.font.name='Consolas'; r.font.size=Pt(8.5); r.font.color.rgb=NAVY
p('การรัน AI ต้องตั้ง KKU API key ใน environment ตาม `.env.example`; ห้ามใส่ secret ใน repository. ตัวเลขในเล่มอ้างอิง snapshot วันที่ 26 กันยายน 2569.')

# References
page(); heading('เอกสารอ้างอิง')
apa_refs=[
    ('Arcuri, A. (2018). Test suite generation with the Many Independent Objective (MIO) algorithm. ', 'Information and Software Technology, 104', ', 195–206. https://doi.org/10.1016/j.infsof.2018.05.003'),
    ('Defects4J contributors. (n.d.). ', 'Defects4J (Version 3.0.1-7-g8c16da82)', ' [Computer software]. GitHub. https://github.com/rjust/defects4j/tree/8c16da8230843cdc918eaf4ddb449637f02b83c6'),
    ('Fraser, G., & Arcuri, A. (2011). EvoSuite: Automatic test suite generation for object-oriented software. In ', 'Proceedings of the 19th ACM SIGSOFT Symposium on the Foundations of Software Engineering', ' (pp. 416–419). https://doi.org/10.1145/2025113.2025179'),
    ('Just, R., Jalali, D., & Ernst, M. D. (2014). Defects4J: A database of existing faults to enable controlled testing studies for Java programs. In ', 'Proceedings of the 2014 International Symposium on Software Testing and Analysis', ' (pp. 437–440). https://doi.org/10.1145/2610384.2628055'),
    ('Lei, Y., & Tai, K. C. (1998). In-parameter-order: A test generation strategy for pairwise testing. In ', 'Proceedings of the 3rd IEEE High-Assurance Systems Engineering Symposium', ' (pp. 254–261). https://doi.org/10.1109/HASE.1998.731623'),
]
for prefix,italic_text,suffix in apa_refs:
    q=DOC.add_paragraph(); q.paragraph_format.left_indent=Inches(.5); q.paragraph_format.first_line_indent=Inches(-.5)
    q.paragraph_format.line_spacing=2.0; q.paragraph_format.space_after=Pt(0)
    q.add_run(prefix); q.add_run(italic_text).italic=True; q.add_run(suffix)
p('รายละเอียดค่าจำนวนบั๊ก, configuration, prompt และผลวัดในรายงานมาจาก catalog, source, manifest, CSV และ run logs ของโครงงาน ณ snapshot ที่ระบุ.')

# Appendices
page(); heading('ภาคผนวก ก แหล่งข้อมูลและไฟล์ทำซ้ำ')
table(['เนื้อหา','ไฟล์หลัก'],[
    ('โจทย์รายวิชา','SQA_Project_2026 (1).pdf หน้า 1–3'),
    ('Catalog','target_benchmark/all_bugs_catalog.json; target_benchmark/catalog_17_projects.json'),
    ('Benchmark results','results/master_benchmark_summary.csv; results/run_logs/'),
    ('Suite gap audit','results/suite_gap_audit.csv; scripts/audit_suite_gaps.py'),
    ('สถิติและ analytics','results/master_descriptive_stats.json; results/advanced_analytics_report.md'),
    ('Native IPO','Combinatorial_IPO/Code/; Combinatorial_IPO/Results/verified_suites_manifest.json; generation_manifest.json; routing_manifest.json'),
    ('MIO','MIO_Algorithm/Code/batch_evosuite.py; MIO_Algorithm/Result_Round2/evosuite_budget_summary.csv; MIO_FAILURE_ANALYSIS_REPORT.md'),
    ('AI prompts/economics','Deepseek-v4_flash/Prompt/; Gemini-3_8_flash/Prompt/; results/Deepseek_vs_Gemini_Economics.csv'),
    ('Runner/analytics','scripts/run_benchmark.py; scripts/run_member3_delivery.py; scripts/consolidate_master_results.py; scripts/advanced_data_analytics.py; results/member3_delivery_logs/')
],[1.65,5.05],'ไฟล์หลักสำหรับตรวจสอบและทำซ้ำ',8.3)
heading('ภาคผนวก ข ตัวอย่างการตรวจย้อนกลับ',1)
samples=[]
for t in TECHS:
    r=next(x for x in ROWS if x['Technique']==t and x['Fault_Detection_Status']=='BUG_DETECTED')
    samples.append((LABEL[t],f"{r['Project']}-{r['Bug_ID']}",r['Run_ID'],r['Test_Files'],r['Suite_SHA256'][:16]+'…'))
table(['เทคนิค','Bug','Run ID','Test file','SHA-256 prefix'],samples,[1.25,.72,1.48,2.0,1.25],'ตัวอย่าง BUG_DETECTED ที่ตรวจ buggy/fixed และ provenance แล้ว',7.4)
p('ทั้งสี่ตัวอย่างมี buggy failures และไม่มี fixed failures ใน structured run log. Hash เต็มและรายการ failure อยู่ใน Run_Log ที่อ้างด้วย project + bug + technique ใน master CSV.')
heading('ภาคผนวก ค Data dictionary แบบย่อ',1)
table(['ฟิลด์','ความหมาย'],[
    ('Project, Bug_ID, Technique','คีย์หนึ่งการประเมินต่อบั๊กและเทคนิค'),
    ('Target_Classes','modified target classes ในขอบเขตบั๊ก'),
    ('Suite_Available, Test_Files','การมี suite และชื่อไฟล์ที่ประเมิน'),
    ('Line_Coverage_%, Branch_Coverage_%','ค่าจาก Defects4J summary; ว่างเมื่อไม่ได้วัด'),
    ('Fault_Detection_Status','BUG_DETECTED, NOT_DETECTED, FLAKY_OR_REGRESSION, COMPILE_ERROR หรือ NOT_EVALUATED'),
    ('Execution_Status','สถานะ benchmark; สถานะ generation อยู่ใน audit; hash และ log อยู่ใน master')
],[2.25,4.45],'ความหมายฟิลด์สำหรับตรวจย้อนกลับ',8.5)
tail=DOC.paragraphs[-1]
tail.paragraph_format.space_before=Pt(0); tail.paragraph_format.space_after=Pt(0)
tail.paragraph_format.line_spacing=Pt(1)
for run in tail.runs: run.font.size=Pt(1)

# Footer and metadata
f=SEC.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('CP353201 Software Quality Assurance  |  '); r.font.name=FONT; r.font.size=Pt(8); r.font.color.rgb=GRAY
field(f,' PAGE ','2')
DOC.core_properties.title='รายงานฉบับสมบูรณ์ การเปรียบเทียบการสร้างชุดทดสอบบน Defects4J'
DOC.core_properties.subject='CP353201 Software Quality Assurance ภาคการศึกษา 1/2569'
DOC.core_properties.author='กลุ่มโครงงาน CP353201'
up=OxmlElement('w:updateFields'); up.set(qn('w:val'),'true'); DOC.settings.element.append(up)
DOC.save(OUT)
print('created',OUT,'bytes',OUT.stat().st_size)
